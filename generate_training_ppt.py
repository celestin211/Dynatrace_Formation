from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE


TITLE_COLOR = RGBColor(0x14, 0x4A, 0x7C)
TEXT_COLOR = RGBColor(0x1F, 0x1F, 0x1F)
ACCENT_COLOR = RGBColor(0x00, 0x84, 0xD1)
LIGHT_BG = RGBColor(0xF3, 0xF8, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DYNATRACE_GREEN = RGBColor(0x6C, 0xD6, 0x4D)
DYNATRACE_BLUE = RGBColor(0x14, 0x8E, 0xF3)
DARK_BG = RGBColor(0x0F, 0x17, 0x2A)


def set_title_style(shape):
    p = shape.text_frame.paragraphs[0]
    if p.runs:
        run = p.runs[0]
        run.font.bold = True
        run.font.size = Pt(34)
        run.font.color.rgb = TITLE_COLOR


def style_slide_background(slide, dark=False):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(13.33),
        Inches(7.5),
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = TITLE_COLOR if dark else LIGHT_BG
    bg.line.fill.background()
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)


def add_footer(slide, text):
    footer = slide.shapes.add_textbox(Inches(0.4), Inches(7.05), Inches(12.5), Inches(0.3))
    tf = footer.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    for run in p.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)


def add_brand_tag(slide, text="Dynatrace Formation"):
    tag = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(10.7),
        Inches(0.25),
        Inches(2.3),
        Inches(0.45),
    )
    tag.fill.solid()
    tag.fill.fore_color.rgb = DYNATRACE_GREEN
    tag.line.fill.background()
    tf = tag.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    for run in p.runs:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = DARK_BG


def add_divider(slide):
    div = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.5),
        Inches(1.2),
        Inches(12.3),
        Inches(0.08),
    )
    div.fill.solid()
    div.fill.fore_color.rgb = DYNATRACE_BLUE
    div.line.fill.background()


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    style_slide_background(slide, dark=True)
    slide.shapes.title.text = title
    set_title_style(slide.shapes.title)
    for run in slide.shapes.title.text_frame.paragraphs[0].runs:
        run.font.color.rgb = WHITE
    slide.placeholders[1].text = subtitle
    tf = slide.placeholders[1].text_frame
    for p in tf.paragraphs:
        for run in p.runs:
            run.font.size = Pt(18)
            run.font.color.rgb = WHITE
    add_brand_tag(slide, "Client Ready Deck")
    add_footer(slide, "Formation Dynatrace | Support visuel v3")


def add_bullet_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    style_slide_background(slide)
    slide.shapes.title.text = title
    set_title_style(slide.shapes.title)
    add_brand_tag(slide)
    add_divider(slide)
    body = slide.shapes.placeholders[1].text_frame
    body.clear()

    for idx, bullet in enumerate(bullets):
        p = body.paragraphs[0] if idx == 0 else body.add_paragraph()
        if isinstance(bullet, tuple):
            text, level = bullet
        else:
            text, level = bullet, 0
        p.text = text
        p.level = level
        for run in p.runs:
            run.font.size = Pt(22 if level == 0 else 18)
            run.font.color.rgb = TEXT_COLOR
    add_footer(slide, "Dynatrace Formation 2026")


def add_section_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    style_slide_background(slide)
    hero = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(13.33),
        Inches(7.5),
    )
    hero.fill.solid()
    hero.fill.fore_color.rgb = DARK_BG
    hero.line.fill.background()
    slide.shapes._spTree.remove(hero._element)
    slide.shapes._spTree.insert(3, hero._element)
    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(11.8), Inches(1.5))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    for run in p.runs:
        run.font.size = Pt(44)
        run.font.bold = True
        run.font.color.rgb = WHITE
    sp = tf.add_paragraph()
    sp.text = subtitle
    for run in sp.runs:
        run.font.size = Pt(20)
        run.font.color.rgb = WHITE
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.9),
        Inches(4.55),
        Inches(3.7),
        Inches(0.12),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = DYNATRACE_GREEN
    accent.line.fill.background()
    add_footer(slide, "Section")


def add_two_column_slide(prs, title, left_title, left_items, right_title, right_items):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    style_slide_background(slide)
    slide.shapes.title.text = title
    set_title_style(slide.shapes.title)
    add_brand_tag(slide)
    add_divider(slide)

    left_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(5.7), Inches(5.4))
    right_box = slide.shapes.add_textbox(Inches(6.1), Inches(1.5), Inches(6.8), Inches(5.4))

    left_tf = left_box.text_frame
    left_tf.word_wrap = True
    left_tf.clear()
    p = left_tf.paragraphs[0]
    p.text = left_title
    for run in p.runs:
        run.font.bold = True
        run.font.size = Pt(24)
        run.font.color.rgb = ACCENT_COLOR
    for item in left_items:
        bp = left_tf.add_paragraph()
        bp.text = f"- {item}"
        bp.level = 0
        for run in bp.runs:
            run.font.size = Pt(18)
            run.font.color.rgb = TEXT_COLOR

    right_tf = right_box.text_frame
    right_tf.word_wrap = True
    right_tf.clear()
    p = right_tf.paragraphs[0]
    p.text = right_title
    for run in p.runs:
        run.font.bold = True
        run.font.size = Pt(24)
        run.font.color.rgb = ACCENT_COLOR
    for item in right_items:
        bp = right_tf.add_paragraph()
        bp.text = f"- {item}"
        bp.level = 0
        for run in bp.runs:
            run.font.size = Pt(18)
            run.font.color.rgb = TEXT_COLOR
    add_footer(slide, "Comparatif")


def add_client_cover_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    style_slide_background(slide, dark=True)
    ribbon = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(13.33),
        Inches(0.35),
    )
    ribbon.fill.solid()
    ribbon.fill.fore_color.rgb = DYNATRACE_GREEN
    ribbon.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.8), Inches(11.8), Inches(2.0))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "Formation Dynatrace"
    for run in p.runs:
        run.font.size = Pt(52)
        run.font.bold = True
        run.font.color.rgb = WHITE

    p2 = tf.add_paragraph()
    p2.text = "OneAgent • Kubernetes • DQL • Dashboards"
    for run in p2.runs:
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(0xD6, 0xEE, 0xFF)

    p3 = tf.add_paragraph()
    p3.text = "Durée 3h | Niveau débutant/intermédiaire"
    for run in p3.runs:
        run.font.size = Pt(18)
        run.font.color.rgb = WHITE

    badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.9),
        Inches(5.9),
        Inches(3.3),
        Inches(0.55),
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = DYNATRACE_BLUE
    badge.line.fill.background()
    btf = badge.text_frame
    btf.clear()
    bp = btf.paragraphs[0]
    bp.text = "Session client - v3"
    for run in bp.runs:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = WHITE

    add_footer(slide, "À personnaliser : logo client / date / formateur")


def add_module_agenda_slide(prs):
    add_bullet_slide(
        prs,
        "Programme de la formation (3h)",
        [
            "Module 1 - Accueil & Introduction (25 min)",
            "Module 2 - Naviguer dans Dynatrace (50 min)",
            "Module 3 - Pause (15 min)",
            "Module 4 - Logs, Métriques & Alertes (45 min)",
            "Module 5 - Créer un Dashboard (35 min)",
            "Module 6 - Bilan & Questions (10 min)",
        ],
    )


def build_presentation(output_path):
    prs = Presentation()

    add_client_cover_slide(prs)

    add_title_slide(
        prs,
        "Connaître Dynatrace en 3 Heures",
        "OneAgent | Kubernetes | Cloud | Dashboards\nNiveau débutant | Version mai 2026",
    )

    add_module_agenda_slide(prs)
    add_section_slide(prs, "Module 1 & 2", "Fondamentaux, navigation et incidents")

    add_bullet_slide(
        prs,
        "Objectifs pédagogiques",
        [
            "Naviguer dans l'interface Dynatrace",
            "Analyser logs, métriques, traces et topologie",
            "Comprendre les incidents détectés par Davis AI",
            "Créer un dashboard de supervision personnalisable",
            "Exécuter un mini atelier Kubernetes + Dynatrace",
        ],
    )

    add_two_column_slide(
        prs,
        "Dynatrace en bref",
        "Sans Dynatrace",
        [
            "Incidents détectés tardivement",
            "Analyse manuelle sur plusieurs outils",
            "Cause racine difficile à identifier",
        ],
        "Avec Dynatrace",
        [
            "Détection proactive avec Davis AI",
            "Données unifiées dans une seule plateforme",
            "Topologie et dépendances détectées automatiquement",
        ],
    )

    add_bullet_slide(
        prs,
        "Les 4 types de données",
        [
            "Métriques : CPU, mémoire, latence, débit",
            "Logs : événements, erreurs, debug",
            "Traces : parcours bout-en-bout d'une requête",
            "Topologie : carte des relations entre services et hôtes",
        ],
    )

    add_bullet_slide(
        prs,
        "OneAgent: composant essentiel",
        [
            "Installé sur hôtes, VM ou conteneurs",
            "Auto-discovery des processus, services et dépendances",
            "Collecte métriques, traces et logs (selon config)",
            "Améliore la précision de l'analyse Davis AI",
        ],
    )

    add_bullet_slide(
        prs,
        "Module 2 - Navigation dans Dynatrace",
        [
            "Infrastructure -> Hosts : santé des hôtes",
            "Applications -> Services : erreurs, latence, débit",
            "Infrastructure -> Smartscape: topologie applicative",
            "Alerts -> Problems : incidents et cause racine",
        ],
    )

    add_bullet_slide(
        prs,
        "Exercice guide - Analyse d'un Problem",
        [
            "Ouvrir Alerts -> Problems",
            "Sélectionner un incident ouvert ou résolu",
            "Identifier les entités impactées",
            "Lire la timeline et la root cause Davis AI",
            "Valider le plan d'action de remédiation",
        ],
    )

    add_section_slide(prs, "Module 4 & 5", "Logs, DQL et dashboards")

    add_bullet_slide(
        prs,
        "Module 4 - Logs et DQL",
        [
            "Filtrer rapidement les erreurs: loglevel == \"ERROR\"",
            "Identifier les services les plus bruyants",
            "Visualiser l'évolution des erreurs dans le temps",
            "Transformer une requête DQL en tile dashboard",
        ],
    )

    add_bullet_slide(
        prs,
        "5 requêtes DQL indispensables",
        [
            "Compter les erreurs globales",
            "Lister les dernières erreurs avec contexte",
            "Classer les erreurs par service",
            "Tracer les erreurs par tranche de 5 minutes",
            "Filtrer dynamiquement avec une variable {{$service}}",
        ],
    )

    add_bullet_slide(
        prs,
        "Module 5 - Construire un Dashboard",
        [
            "Tile KPI : nombre d'erreurs",
            "Tile Line chart : évolution des erreurs",
            "Tile Table : Top 5 services en erreur",
            "Variables dynamiques : $service, $env",
            "Partage du dashboard avec l'équipe",
        ],
    )

    add_bullet_slide(
        prs,
        "Atelier Kubernetes + Dynatrace",
        [
            "Prérequis : cluster Kubernetes + Dynatrace Operator",
            "Déploiement namespace, deployment, service, ingress",
            "Injection OneAgent via annotations Kubernetes",
            "Vérification des pods et de la visibilité Dynatrace",
            "Test d'accès via Ingress NGINX (web-demo.local)",
        ],
    )

    add_bullet_slide(
        prs,
        "Commandes atelier (résumé)",
        [
            "kubectl apply -f k8s/namespace.yaml",
            "kubectl apply -f k8s/deployment.yaml",
            "kubectl apply -f k8s/service.yaml",
            "kubectl apply -f k8s/ingress.yaml",
            "kubectl get pods,svc,ingress -n dynatrace-demo",
        ],
    )

    add_section_slide(prs, "Module 6", "Bilan, quiz et prochaines étapes")

    add_bullet_slide(
        prs,
        "Bonnes pratiques",
        [
            "Ne jamais versionner des binaires lourds ou des tokens",
            "Commencer avec des seuils d'alerte larges puis affiner",
            "Construire des dashboards par usage (ops/dev/management)",
            "Documenter les requêtes DQL réutilisables",
            "Sécuriser les accès et privilégier le moindre privilège",
        ],
    )

    add_bullet_slide(
        prs,
        "Bilan & quiz final",
        [
            "Rôle de Davis AI dans la détection et la root cause",
            "Filtrage de logs avec DQL",
            "Différence Problem vs alerte de seuil",
            "Création de variable dashboard et usage dans DQL",
            "Plan d'amélioration continue post-formation",
        ],
    )

    add_bullet_slide(
        prs,
        "Ressources utiles",
        [
            "Dynatrace University: university.dynatrace.com",
            "Documentation: docs.dynatrace.com",
            "Community: community.dynatrace.com",
            "Parcours recommandé : certification Associate",
        ],
    )

    add_title_slide(
        prs,
        "Merci !",
        "Questions / Réponses\nContact formateur : à personnaliser",
    )

    prs.save(output_path)


if __name__ == "__main__":
    build_presentation("Dynatrace_Formation_3h_v3.pptx")
